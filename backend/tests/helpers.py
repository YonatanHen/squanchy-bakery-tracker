import csv
import io
from pathlib import Path

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.models import Branch, Fridge, Logger, Metric

SAMPLE_FILE = Path(__file__).resolve().parents[2] / "data" / "sample_week.xlsx"

# What the user confirms in the "add it?" dialog for the sample file's branches and loggers.
SAMPLE_REGISTRATION = {
    "branches": [{"name": name, "city": name} for name in ("Jerusalem", "Tel Aviv", "Haifa", "Rishon LeZion")],
    "fridges": [
        {"logger_id": "TL-0512", "branch": "Jerusalem", "fridge": "Dairy"},
        {"logger_id": "TL-0417", "branch": "Tel Aviv", "fridge": "Walk-in"},
        {"logger_id": "TL-0231", "branch": "Haifa", "fridge": "Dairy", "metric": "F"},
        {"logger_id": "TL-0388", "branch": "Rishon LeZion", "fridge": "Cream cakes"},
    ],
}


def xlsx_bytes(header, rows) -> bytes:
    """Build an .xlsx file in memory with a header row and data rows."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(header))
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def sample_as_csv() -> bytes:
    """Return the rows of data/sample_week.xlsx as a comma-separated CSV file."""
    buffer = io.StringIO(newline="")
    csv.writer(buffer).writerows(load_workbook(SAMPLE_FILE, read_only=True).active.iter_rows(values_only=True))
    return buffer.getvalue().encode("utf-8")


def load_sample(session, register=SAMPLE_REGISTRATION):
    """Ingest the real data/sample_week.xlsx the way an upload does, with the user's confirmed registration."""
    from app.services.ingest.service import ingest_file

    return ingest_file(session, SAMPLE_FILE.name, SAMPLE_FILE.read_bytes(), register)


def add_fridge(session, branch="Jerusalem", fridge="Dairy", logger="TL-0512", metric=Metric.C) -> Fridge:
    """Create a fridge (and its branch if missing) with one logger."""
    existing = session.scalar(select(Branch).where(func.lower(Branch.name) == branch.lower()))
    new_fridge = Fridge(branch=existing or Branch(name=branch, city=branch), name=fridge, metric=metric)
    if logger:
        new_fridge.logger = Logger(id=logger)
    session.add(new_fridge)
    session.commit()
    return new_fridge
